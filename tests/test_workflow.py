"""ワークフロー本体（軽量版）の自己テスト。  実行: python3 -m unittest discover -s tests -v"""
from __future__ import annotations

import contextlib
import importlib.machinery
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path

os.environ["PYTHONDONTWRITEBYTECODE"] = "1"

# ~/.claude をテスト用の一時ディレクトリにする（実環境の ~/.claude に触れない）
TEST_HOME = Path(tempfile.mkdtemp(prefix="sdlc-home-"))
os.environ["HOME"] = str(TEST_HOME)
os.environ["USERPROFILE"] = str(TEST_HOME)   # Windows の Python は HOME ではなく USERPROFILE を見る
os.environ.pop("CLAUDECODE", None)           # Claude Code の中で実行されても approve の検査を再現できるように

REPO = Path(__file__).resolve().parents[1]
BIN = REPO / ".claude" / "sdlc" / "bin"
HOOKS = REPO / ".claude" / "hooks" / "sdlc"
INSTALLER = REPO / "install.py"
sys.path.insert(0, str(REPO / ".claude" / "sdlc" / "lib"))

from sdlc_lib.shell import find_bash  # noqa: E402

BASH = find_bash() or "bash"

from sdlc_lib import checks, fm, gates, items, records, state, testitems, trace  # noqa: E402
from sdlc_lib.config import gated, load_project, phase_of_path, project_problems  # noqa: E402


def load_script(name):
    loader = importlib.machinery.SourceFileLoader(name, str(BIN / name))
    spec = importlib.util.spec_from_loader(name, loader)
    mod = importlib.util.module_from_spec(spec)
    loader.exec_module(mod)
    return mod


approve = load_script("approve")
sdlc_cli = load_script("sdlc")

PROJECT_YAML = """\
workflow: sdlc
language: python
test_framework: unittest
commands:
  test: "python3 -m unittest discover -s tests -t . -q"
src_dir: src
test_dir: tests
"""

SPEC_DOC = """\
# 要求

### REQ-001: 足し算
- 追跡元: SRC-001

# 仕様

### SPEC-001: 足し算の画面
- 追跡元: REQ-001
- 振る舞い: 2 つの数を入れると和を表示する

### NFR-001: 応答
- 追跡元: REQ-001
- 基準値: 1 秒以内

# テスト項目

### TC-001: 和が表示される
- 追跡元: SPEC-001
- 種別: 手動
- 前提: 画面を開いている
- 手順:
  1. 1 と 2 を入れる
  2. 計算を押す
- 確認ポイント: 結果欄
- 期待結果: 3 と表示される

### TC-002: add の計算
- 追跡元: SPEC-001
- 種別: 自動
- 期待結果: add(1, 2) == 3

### TC-003: 応答時間
- 追跡元: NFR-001
- 種別: 両方
- 期待結果: 1 秒以内に表示される
"""

TEST_CODE = """\
# trace: TC-002, TC-003
import unittest
from src.calc import add


class T(unittest.TestCase):
    def test_add(self):
        self.assertEqual(add(1, 2), 3)
"""


class Repo:
    """一時 git リポジトリ上のサンプルプロジェクト。"""

    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="sdlc-test-"))
        self.run("git", "init", "-q")
        self.run("git", "config", "user.email", "t@example.com")
        self.run("git", "config", "user.name", "tester")
        self.write("project.yaml", PROJECT_YAML)
        self.write(".gitignore", "__pycache__/\n.sdlc/\n")
        for d in ("src", "tests"):
            self.write(f"{d}/__init__.py", "")
        self.write("input/req.md", "# 要求\n\n- 足し算ができること\n")
        state.save(self.dir, dict(state.DEFAULT_STATE))

    def run(self, *cmd, **kw):
        return subprocess.run(cmd, cwd=self.dir, capture_output=True, text=True, encoding="utf-8", errors="replace", **kw)

    def write(self, rel, text):
        p = self.dir / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")

    @property
    def proj(self):
        return load_project(self.dir)

    def approve(self, phase, decision="approve"):
        a = Namespace(target=phase, decision=decision, approver="tester", comment="ok", yes=True)
        with contextlib.redirect_stdout(io.StringIO()):
            approve.approve_phase(self.dir, self.proj, a)

    def hook(self, script, tool_input, env=None, tool_name=None):
        data = {"tool_input": tool_input, "cwd": str(self.dir)}
        if tool_name:
            data["tool_name"] = tool_name
        e = dict(os.environ, CLAUDE_PROJECT_DIR=str(self.dir))
        e.pop("SDLC_MAINTENANCE", None)
        e.update(env or {})
        return subprocess.run([sys.executable, str(HOOKS / script)], input=json.dumps(data), capture_output=True,
                              text=True, encoding="utf-8", errors="replace", cwd=self.dir, env=e)

    def can_write(self, rel, **kw):
        r = self.hook("guard_write.py", {"file_path": str(self.dir / rel), "content": "x"}, **kw)
        return r.returncode == 0, r.stderr

    def sdlc(self, *args):
        """scripts/sdlc 相当を同じプロセスで実行し、(終了コード, 出力) を返す。"""
        out, err = io.StringIO(), io.StringIO()
        code = 0
        old = os.getcwd()
        os.chdir(self.dir)
        try:
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                sdlc_cli.main(list(args))
        except SystemExit as e:
            code = e.code if isinstance(e.code, int) else 1
        finally:
            os.chdir(old)
        return code, out.getvalue() + err.getvalue()

    def cleanup(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class Base(unittest.TestCase):
    def setUp(self):
        self.r = Repo()
        self._cwd = os.getcwd()
        os.chdir(self.r.dir)

    def tearDown(self):
        os.chdir(self._cwd)
        self.r.cleanup()

    def spec(self, approve_it=True):
        self.r.write("docs/01-specification/intake/req.md", "### SRC-001\n\n> 足し算\n")
        self.r.write("docs/01-specification/spec.md", SPEC_DOC)
        if approve_it:
            self.r.approve("01")

    def implement(self):
        self.r.write("docs/02-design/design.md", "# 設計\n\n### DS-001: calc\n- 追跡元: SPEC-001\n")
        self.r.write("src/calc.py", "# trace: DS-001\ndef add(a, b):\n    return a + b\n")
        self.r.write("tests/test_calc.py", TEST_CODE)

    def review_ok(self, phase):
        self.r.write(f"docs/{phase}/review.md", "# レビュー\n\n- 判定: 合格\n- 残指摘: 0\n")

    def fill_results(self, result="OK", tcs=("TC-001", "TC-003")):
        p = testitems.results_path(self.r.dir)
        text = p.read_text(encoding="utf-8")
        for tc in tcs:
            lines = []
            for line in text.splitlines():
                if line.startswith(f"| {tc} |"):
                    cells = line.strip("|").split("|")
                    cells[3], cells[4], cells[5] = f" {result} ", " 山田 ", " 2026-10-29 "
                    line = "|" + "|".join(cells) + "|"
                lines.append(line)
            text = "\n".join(lines) + "\n"
        p.write_text(text, encoding="utf-8")


class TestConfig(Base):
    def test_project_and_areas(self):
        p = self.r.proj
        self.assertEqual(project_problems(p), [])
        self.assertEqual(p["gates"], ["01-specification"])
        self.assertEqual(phase_of_path("input/x.xlsx", p), "01-specification")
        self.assertEqual(phase_of_path("docs/02-design/d.md", p), "02-design")
        self.assertEqual(phase_of_path("src/a.py", p), "03-implementation")
        self.assertEqual(phase_of_path("tests/test_a.py", p), "03-implementation")
        self.assertEqual(phase_of_path("README.md", p), None)

    def test_gates_option(self):
        self.r.write("project.yaml", PROJECT_YAML + "gates: [02-design, 4]\n")
        p = self.r.proj
        self.assertEqual(p["gates"], ["01-specification", "02-design", "04-acceptance"])
        self.assertTrue(gated(p, "02") and not gated(p, "03"))
        self.r.write("project.yaml", PROJECT_YAML + "gates: []\n")
        self.assertEqual(self.r.proj["gates"], ["01-specification"])   # 01 は外せない
        self.r.write("project.yaml", PROJECT_YAML + "gates: [99-x]\n")
        self.assertTrue(any("gates" in x for x in project_problems(self.r.proj)))

    def test_dirs(self):
        self.r.write("project.yaml", PROJECT_YAML.replace("src_dir: src", "src_dir: docs/src"))
        self.assertTrue(any("予約" in x for x in project_problems(self.r.proj)))
        self.r.write("project.yaml", PROJECT_YAML.replace("test_dir: tests", "test_dir: src/tests"))
        self.assertTrue(any("重ならない" in x for x in project_problems(self.r.proj)))


class TestTrace(Base):
    def test_chain_and_tc_kind(self):
        self.spec(approve_it=False)
        g = trace.build(self.r.dir, self.r.proj)
        self.assertEqual(trace.check(g, 0), [])
        self.assertEqual(trace.tc_kind(g, "TC-003"), "両方")
        # 03 からは自動テストの TC にテストコードが要る
        issues = trace.check(g, 2)
        self.assertEqual(sorted(i.split(":")[0] for i in issues),
                         ["[自動テストなし] TC-002", "[自動テストなし] TC-003"])
        self.implement()
        g = trace.build(self.r.dir, self.r.proj)
        self.assertEqual(trace.check(g, 2), [])
        self.assertEqual(trace.automated(g, "TC-002"), ["tests/test_calc.py"])
        imp = trace.impact(g, ["SPEC-001"])
        self.assertIn("TC-001", imp["tcs"])
        self.assertIn("src/calc.py", imp["files"])   # 設計 DS 経由のコードも影響範囲に入る

    def test_missing_links(self):
        self.r.write("docs/01-specification/spec.md",
                     "### REQ-001: a\n- 追跡元: SRC-404\n\n### SPEC-001: b\n- 追跡元: REQ-001\n\n"
                     "### TC-001: c\n- 追跡元: SPEC-001\n\n### SPEC-002: d\n")
        issues = trace.check(trace.build(self.r.dir, self.r.proj), 0)
        text = "\n".join(issues)
        self.assertIn("[未定義の参照] SRC-404", text)
        self.assertIn("[種別なし] TC-001", text)
        self.assertIn("[追跡元なし] SPEC-002", text)
        self.assertIn("[下流なし] SPEC-002", text)

    def test_exempt(self):
        self.r.write("docs/01-specification/spec.md",
                     "- 除外: SRC-002（目次）\n\n### SRC-002: 目次\n\n### Q-001: 質問\n\n"
                     "### NFR-001: 国際化\n- 追跡元: Q-001\n- 対象外: 日本語のみ\n")
        self.assertEqual(trace.check(trace.build(self.r.dir, self.r.proj), 0), [])


class TestGates(Base):
    def test_only_gated_phases(self):
        self.spec()
        self.assertTrue(gates.gate_status(self.r.dir, self.r.proj, "01")["valid"])
        self.r.write("docs/02-design/design.md", "# 設計\n")
        with self.assertRaises(SystemExit):
            self.r.approve("02")       # 人間ゲートの無い工程は承認しない
        self.r.write("docs/01-specification/spec.md", SPEC_DOC + "\n追記\n")
        st = gates.gate_status(self.r.dir, self.r.proj, "01")
        self.assertFalse(st["valid"])
        self.assertIn("ハッシュ不一致", st["reason"])

    def test_optional_gate(self):
        self.r.write("project.yaml", PROJECT_YAML + "gates: [02-design]\n")
        self.spec()
        self.r.write("docs/02-design/design.md", "# 設計\n")
        self.review_ok("02-design")
        self.r.approve("02")
        self.assertTrue(gates.gate_status(self.r.dir, self.r.proj, "02")["valid"])
        self.assertFalse(self.r.can_write("docs/02-design/design.md")[0])

    def test_blockers(self):
        self.spec(approve_it=False)
        self.r.write("docs/01-specification/review.md", "# レビュー\n\n- 判定: 不合格\n")
        with self.assertRaises(SystemExit):
            self.r.approve("01")
        self.r.approve("01", "reject")
        self.assertEqual(fm.read(gates.gate_path(self.r.dir, "01"))[0]["decision"], "rejected")
        self.review_ok("01-specification")
        self.r.write("docs/01-specification/spec.md", "### REQ-001: x\n- 追跡元: SRC-404\n")
        with self.assertRaises(SystemExit):
            self.r.approve("01")       # トレースの MUST 失敗

    def test_history_kept(self):
        self.spec()
        self.r.write("docs/01-specification/questions.md", "# 質問\n")
        self.r.approve("01")
        _, body = fm.read(gates.gate_path(self.r.dir, "01"))
        self.assertIn("## 履歴", body)


class TestGuardWrite(Base):
    def test_gates_and_results_denied(self):
        self.assertFalse(self.r.can_write("state/gates/01-specification.md")[0])
        self.assertFalse(self.r.can_write("state/gates/x.md", env={"SDLC_MAINTENANCE": "1"})[0])
        self.spec()
        ok, msg = self.r.can_write("docs/04-acceptance/results.md")
        self.assertFalse(ok)
        self.assertIn("人が記入", msg)
        self.assertTrue(self.r.can_write("docs/04-acceptance/plan.md")[0])

    def test_protected_paths(self):
        self.assertFalse(self.r.can_write(".claude/settings.json")[0])
        self.assertFalse(self.r.can_write(".claude/sdlc/lib/sdlc_lib/gates.py")[0])
        self.assertFalse(self.r.can_write(".claude/hooks/sdlc/guard_write.py")[0])
        self.assertFalse(self.r.can_write("scripts/sdlc")[0])
        self.assertFalse(self.r.can_write(".claude/skills/phase-design/SKILL.md")[0])
        self.assertFalse(self.r.can_write(".claude/agents/tester.md")[0])
        self.assertTrue(self.r.can_write(".claude/skills/my-own-skill/SKILL.md")[0])
        self.assertTrue(self.r.can_write(".claude/settings.json", env={"SDLC_MAINTENANCE": "1"})[0])

    def test_spec_gate_controls_downstream(self):
        self.assertTrue(self.r.can_write("docs/01-specification/spec.md")[0])
        ok, msg = self.r.can_write("docs/02-design/design.md")
        self.assertFalse(ok)
        self.assertIn("01-specification", msg)
        self.assertFalse(self.r.can_write("src/calc.py")[0])
        self.spec()
        # 承認後: 仕様は変更不可、設計・実装・テストは自由に直せる（人間ゲートなし）
        ok, msg = self.r.can_write("docs/01-specification/spec.md")
        self.assertFalse(ok)
        self.assertIn("/change", msg)
        self.assertFalse(self.r.can_write("input/req.md")[0])
        for rel in ("docs/02-design/design.md", "src/calc.py", "tests/test_calc.py", "docs/04-acceptance/plan.md"):
            self.assertTrue(self.r.can_write(rel)[0], rel)
        # TDD ガードは無い（失敗テストが無くても src を編集できる）
        self.implement()
        self.assertTrue(self.r.can_write("src/calc.py")[0])

    def test_free_paths(self):
        self.assertTrue(self.r.can_write("state/defects/DEF-001.md")[0])
        self.assertTrue(self.r.can_write("README.md")[0])
        self.assertTrue(self.r.can_write("rules/common.md")[0])


class TestGuardBashAndSkill(Base):
    def bash(self, cmd, tool="Bash"):
        return self.r.hook("guard_bash.py", {"command": cmd}, tool_name=tool).returncode == 0

    def test_bash(self):
        self.assertTrue(self.bash("cat state/gates/01-specification.md"))
        self.assertFalse(self.bash("echo x > state/gates/01-specification.md"))
        self.assertFalse(self.bash("scripts/approve 01 --yes"))
        self.assertFalse(self.bash("python3 scripts/approve 1"))
        self.assertTrue(self.bash("git commit -m 'gate は scripts/approve で記録'"))
        self.assertFalse(self.bash("echo '{}' > .claude/settings.local.json"))
        self.assertFalse(self.bash("echo x > docs/01-specification/a.md"))
        self.assertFalse(self.bash("cp /tmp/x input/req.md"))
        self.assertFalse(self.bash("sed -i s/NG/OK/ docs/04-acceptance/results.md"))
        self.assertFalse(self.bash("cp /tmp/r.xlsx docs/04-acceptance/results.xlsx"))
        self.assertTrue(self.bash("sed -i s/a/b/ src/calc.py"))        # 人間ゲートの無い工程は検査しない
        self.assertTrue(self.bash("echo x > docs/02-design/notes.md"))
        self.assertFalse(self.bash("cp /tmp/x .claude/hooks/sdlc/guard_write.py"))
        self.assertFalse(self.bash("python3 install.py --user"))
        self.assertFalse(self.bash("echo x > ~/.claude/sdlc/RULES.md"))
        self.assertTrue(self.bash("python3 -m unittest discover -s tests 2>&1 | tail"))
        self.assertTrue(self.bash("scripts/sdlc testitems sheet"))
        self.assertTrue(self.bash("scripts/sdlc testitems import ~/Downloads/results.xlsx"))

    def test_bash_with_optional_gate(self):
        self.r.write("project.yaml", PROJECT_YAML + "gates: [03-implementation]\n")
        self.assertFalse(self.bash("sed -i s/a/b/ src/calc.py"))

    def test_powershell(self):
        ps = "PowerShell"
        self.assertFalse(self.bash("Set-Content docs\\01-specification\\a.md 'x'", ps))
        self.assertFalse(self.bash("& .\\scripts\\approve.cmd 01", ps))
        self.assertFalse(self.bash("'x' | Out-File docs\\04-acceptance\\results.md", ps))
        self.assertTrue(self.bash("scripts\\sdlc.cmd status", ps))

    def test_skill(self):
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "phase-specification"}).returncode, 0)
        r = self.r.hook("guard_skill.py", {"skill": "phase-design"})
        self.assertEqual(r.returncode, 2)
        self.spec()
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "phase-design"}).returncode, 0)
        self.assertEqual(self.r.hook("guard_skill.py", {"skill": "phase-acceptance"}).returncode, 0)


class TestChecks(Base):
    def test_levels_and_skip(self):
        self.r.write("project.yaml", PROJECT_YAML + 'checks:\n  common:\n    must: ["true"]\n    should: ["false"]\n'
                     '    may: ["no_such_tool_xyz --version"]\n')
        res = checks.run(self.r.dir, self.r.proj, "01")
        self.assertTrue(res["must_passed"])
        self.assertEqual(res["counts"]["should"]["fail"], 1)
        self.assertEqual(res["skipped"], ["no_such_tool_xyz"])

    def test_builtin_test_and_acceptance(self):
        self.spec()
        self.implement()
        names = [i["name"] for i in checks.run(self.r.dir, self.r.proj, "02")["items"]]
        self.assertEqual(names, ["trace"])
        res = checks.run(self.r.dir, self.r.proj, "03")
        self.assertTrue(res["must_passed"], checks.render(res))
        self.assertEqual([i["name"] for i in res["items"]], ["trace", "test"])
        self.r.write("src/calc.py", "def add(a, b):\n    return a - b\n")
        self.assertFalse(checks.run(self.r.dir, self.r.proj, "03")["must_passed"])
        self.implement()
        res = checks.run(self.r.dir, self.r.proj, "04")
        acc = [i for i in res["items"] if i["name"] == "acceptance"][0]
        self.assertEqual(acc["status"], "fail")
        self.assertIn("results.md がありません", acc["output"])


class TestTestItems(Base):
    def setUp(self):
        super().setUp()
        self.spec()
        self.implement()
        self.g = lambda: trace.build(self.r.dir, self.r.proj)

    def test_sheet_and_problems(self):
        r = testitems.sync_sheet(self.r.dir, self.g())
        self.assertEqual(r["added"], ["TC-001", "TC-003"])     # 手動・両方だけ（自動は対象外）
        res = testitems.read_results(self.r.dir)
        self.assertEqual(res["TC-001"]["result"], "未実施")
        self.assertEqual(res["TC-001"]["title"], "和が表示される")
        self.assertEqual(len(testitems.problems(self.r.dir, self.g())), 2)
        self.fill_results("OK", ["TC-001"])
        self.fill_results("NG", ["TC-003"])
        probs = testitems.problems(self.r.dir, self.g())
        self.assertEqual(len(probs), 1)
        self.assertIn("TC-003: NG", probs[0])
        self.fill_results("対象外", ["TC-003"])
        self.assertIn("理由", testitems.problems(self.r.dir, self.g())[0])
        self.fill_results("OK", ["TC-003"])
        self.assertEqual(testitems.problems(self.r.dir, self.g()), [])
        # TC を追加して sheet を再実行しても記入済みの結果は残る
        self.r.write("docs/01-specification/spec2.md", "### TC-004: 追加\n- 追跡元: SPEC-001\n- 種別: 手動\n")
        r = testitems.sync_sheet(self.r.dir, self.g())
        self.assertEqual(r["added"], ["TC-004"])
        res = testitems.read_results(self.r.dir)
        self.assertEqual((res["TC-001"]["result"], res["TC-001"]["by"]), ("OK", "山田"))
        # reset で再確認に戻す
        done = testitems.reset(self.r.dir, self.g(), ["TC-001", "TC-002"], "DEF-001")
        self.assertEqual(done, ["TC-001"])
        res = testitems.read_results(self.r.dir)
        self.assertEqual(res["TC-001"]["result"], "未実施")
        self.assertIn("再確認（DEF-001）", res["TC-001"]["memo"])

    def test_excel_roundtrip(self):
        from openpyxl import load_workbook
        testitems.sync_sheet(self.r.dir, self.g())
        out = self.r.dir / "docs/04-acceptance/results.xlsx"
        self.assertEqual(testitems.export_xlsx(self.r.dir, self.g(), out), 2)
        wb = load_workbook(out)
        ws = wb["受入確認"]
        head = [c.value for c in ws[1]]
        self.assertEqual(head, testitems.XLSX_COLUMNS)
        row = {h: c.value for h, c in zip(head, ws[2])}
        self.assertEqual(row["TC"], "TC-001")
        self.assertIn("1 と 2 を入れる", row["手順"])
        self.assertEqual(row["確認ポイント"], "結果欄")
        col = {h: i + 1 for i, h in enumerate(head)}
        ws.cell(2, col["結果"], "OK")
        ws.cell(2, col["確認者"], "佐藤")
        ws.cell(3, col["結果"], "NG")
        ws.cell(3, col["メモ"], "2 秒かかる\n毎回")
        wb.save(out)
        r = testitems.import_xlsx(self.r.dir, self.g(), out)
        self.assertEqual(r["updated"], ["TC-001", "TC-003"])
        res = testitems.read_results(self.r.dir)
        self.assertEqual((res["TC-001"]["result"], res["TC-001"]["by"]), ("OK", "佐藤"))
        self.assertEqual(res["TC-003"]["memo"], "2 秒かかる<br>毎回")
        # 不正な値があれば何も変えない
        ws.cell(2, col["結果"], "まる")
        wb.save(out)
        before = testitems.results_path(self.r.dir).read_text(encoding="utf-8")
        with self.assertRaises(testitems.ItemsError):
            testitems.import_xlsx(self.r.dir, self.g(), out)
        self.assertEqual(testitems.results_path(self.r.dir).read_text(encoding="utf-8"), before)

    def test_report_and_matrix(self):
        testitems.sync_sheet(self.r.dir, self.g())
        self.fill_results("OK", ["TC-001"])
        rep = testitems.report(self.r.dir, self.g())
        self.assertIn("| SPEC-001 | TC-001 | 手動 | OK | — |", rep)
        self.assertIn("| SPEC-001 | TC-002 | 自動 | — | tests/test_calc.py |", rep)
        m = trace.matrix(self.g(), testitems.read_results(self.r.dir))
        self.assertIn("TC-001: OK", m)


class TestFlow(Base):
    """/sdlc の進み方: 01 だけ人が承認し、02〜04 は機械チェックとレビューで進む。"""

    def status(self):
        return sdlc_cli.compute_status(self.r.dir, self.r.proj)["next"]

    def step(self, s):
        self.assertEqual(self.r.sdlc("phase", "step", s)[0], 0)

    def test_full_flow(self):
        self.assertEqual(self.status()["action"], "run_phase")
        self.spec(approve_it=False)
        self.step("gate")
        self.assertEqual(self.status()["action"], "await_approval")
        code, out = self.r.sdlc("phase", "advance")
        self.assertEqual(code, 1)
        self.assertIn("scripts/approve", out)
        self.r.approve("01")
        self.assertEqual(self.status()["action"], "advance")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        # 02 設計: 人間ゲートなし。レビューが無ければ完了できない
        self.assertEqual(self.status()["skill"], "phase-design")
        self.r.write("docs/02-design/design.md", "# 設計\n")
        self.step("gate")
        self.assertEqual(self.status()["action"], "advance")
        code, out = self.r.sdlc("phase", "advance")
        self.assertEqual(code, 1)
        self.assertIn("review.md", out)
        self.r.write("docs/02-design/review.md", "- 判定: 不合格\n")
        self.assertIn("不合格", self.r.sdlc("phase", "advance")[1])
        self.review_ok("02-design")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        # 03 実装: 自動テストの TC にテストコードが無い・テスト失敗なら完了できない
        self.step("gate")
        code, out = self.r.sdlc("phase", "advance")
        self.assertEqual(code, 1)
        self.assertIn("自動テストなし", out)
        self.implement()
        self.review_ok("03-implementation")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        # 04 受入確認: 人が結果シートに記入するまで完了できない
        self.assertEqual(self.r.sdlc("testitems", "sheet")[0], 0)
        self.step("gate")
        code, out = self.r.sdlc("phase", "advance")
        self.assertEqual(code, 1)
        self.assertIn("TC-001: 未実施", out)
        self.fill_results("OK")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        self.assertEqual(self.status()["action"], "done")
        self.assertTrue(state.load(self.r.dir)["completed"])


class TestItems(Base):
    def to_acceptance(self):
        self.spec()
        self.implement()
        st = state.load(self.r.dir)
        st.update(current_phase="04-acceptance", step="work")
        state.save(self.r.dir, st)
        self.assertEqual(self.r.sdlc("testitems", "sheet")[0], 0)
        self.fill_results("OK")

    def test_defect_cycle(self):
        self.to_acceptance()
        self.fill_results("NG", ["TC-001"])
        code, out = self.r.sdlc("defect", "new", "--title", "和が出ない", "--tc", "TC-001")
        self.assertEqual(code, 0, out)
        self.assertEqual(self.r.sdlc("defect", "start", "DEF-001", "--from", "01")[0], 1)   # 仕様の誤りは /change
        code, out = self.r.sdlc("defect", "start", "DEF-001", "--from", "03")
        self.assertEqual(code, 0, out)
        st = state.load(self.r.dir)
        self.assertEqual((st["mode"], st["current_phase"], st["resume_phase"]), ("item", "03-implementation", "04-acceptance"))
        self.assertEqual(testitems.read_results(self.r.dir)["TC-001"]["result"], "未実施")
        self.assertEqual(sdlc_cli.compute_status(self.r.dir, self.r.proj)["next"]["action"], "run_phase")
        self.review_ok("03-implementation")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        self.assertEqual(sdlc_cli.compute_status(self.r.dir, self.r.proj)["next"]["action"], "item_close")
        code, out = self.r.sdlc("defect", "close", "DEF-001")
        self.assertEqual(code, 1)
        self.assertIn("TC-001: 未実施", out)
        self.fill_results("OK", ["TC-001"])
        code, out = self.r.sdlc("defect", "close", "DEF-001")
        self.assertEqual(code, 0, out)
        st = state.load(self.r.dir)
        self.assertEqual((st["mode"], st["current_phase"]), ("normal", "04-acceptance"))
        self.assertEqual(items.load(self.r.dir, "DEF-001")[0]["status"], "クローズ")

    def test_defect_invalidates_optional_gate(self):
        self.r.write("project.yaml", PROJECT_YAML + "gates: [03-implementation]\n")
        self.spec()
        self.implement()
        self.review_ok("03-implementation")
        self.r.approve("03")
        st = state.load(self.r.dir)
        st.update(current_phase="04-acceptance", step="work")
        state.save(self.r.dir, st)
        self.assertFalse(self.r.can_write("src/calc.py")[0])
        self.r.sdlc("defect", "new", "--title", "x")
        code, out = self.r.sdlc("defect", "start", "DEF-001")
        self.assertEqual(code, 0, out)
        self.assertIn("再承認要", out)
        self.assertTrue(self.r.can_write("src/calc.py")[0])

    def test_change_cycle(self):
        self.to_acceptance()
        self.r.sdlc("change", "new", "--title", "引き算も")
        self.assertEqual(sdlc_cli.compute_status(self.r.dir, self.r.proj)["next"]["action"], "change")
        code, out = self.r.sdlc("change", "start", "CR-001")
        self.assertEqual(code, 0, out)
        self.assertIn("scripts/approve 01-specification", out)
        self.assertFalse(gates.gate_status(self.r.dir, self.r.proj, "01")["valid"])
        self.assertTrue(self.r.can_write("docs/01-specification/spec.md")[0])
        self.assertFalse(self.r.can_write("src/calc.py")[0])          # 仕様の再承認まで下流は止まる
        self.r.write("docs/01-specification/spec.md", SPEC_DOC + "\n### TC-004: 引き算\n- 追跡元: SPEC-001\n- 種別: 手動\n")
        self.r.approve("01")                                           # 唯一の人間ゲート＝変更の承認
        self.assertTrue(self.r.can_write("src/calc.py")[0])
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        self.review_ok("02-design")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        self.review_ok("03-implementation")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)
        self.assertEqual(state.load(self.r.dir)["current_phase"], "04-acceptance")
        self.r.sdlc("testitems", "sheet")
        self.assertEqual(testitems.read_results(self.r.dir)["TC-004"]["result"], "未実施")
        self.assertEqual(self.r.sdlc("change", "close", "CR-001")[0], 0)
        # 新しい TC が未実施なので受入確認は完了できない
        self.r.sdlc("phase", "step", "gate")
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 1)
        self.fill_results("OK", ["TC-004"])
        self.assertEqual(self.r.sdlc("phase", "advance")[0], 0)


class TestSecurity(Base):
    """随時のセキュリティレビュー（/sdlc-security）。工程の状態・完了判定とは独立。"""

    def test_scan_tools(self):
        self.r.write("project.yaml", PROJECT_YAML + 'security:\n  tools: ["true", {name: audit, cmd: "echo CVE-1; false"},'
                     ' "no_such_scanner_xyz"]\n')
        items = checks.security_scan(self.r.dir, self.r.proj)
        self.assertEqual([i["status"] for i in items], ["OK", "検出あり・失敗", "未インストール"])
        text = checks.render_security(items)
        self.assertIn("CVE-1", text)
        code, out = self.r.sdlc("security", "scan", "--out", "docs/security/scan-2026-10-29.md")
        self.assertEqual(code, 0, out)
        self.assertIn("| audit |", (self.r.dir / "docs/security/scan-2026-10-29.md").read_text(encoding="utf-8"))
        self.r.write("project.yaml", PROJECT_YAML)
        self.assertIn("未設定", checks.render_security(checks.security_scan(self.r.dir, self.r.proj)))

    def test_reports_anytime_and_not_blocking(self):
        # どの工程でも（01 の承認前でも）報告を書ける。工程の成果物・ゲートのハッシュには含まれない
        self.assertTrue(self.r.can_write("docs/security/security-review-2026-10-29.md")[0])
        self.assertTrue(self.r.hook("guard_bash.py", {"command": "echo x > docs/security/notes.md"}).returncode == 0)
        self.spec()
        self.r.write("docs/security/security-review-2026-10-29.md",
                     "# セキュリティレビュー報告\n\n- 対象範囲: 全体\n- 対象コミット: abc\n- 所見: Critical 0 / High 1\n")
        self.assertTrue(gates.gate_status(self.r.dir, self.r.proj, "01")["valid"])
        self.assertIsNone(phase_of_path("docs/security/x.md", self.r.proj))
        code, out = self.r.sdlc("security", "list")
        self.assertEqual(code, 0)
        self.assertIn("security-review-2026-10-29.md  対象範囲: 全体 / 対象コミット: abc / 所見: Critical 0 / High 1", out)
        st = state.load(self.r.dir)
        self.assertEqual(st["current_phase"], "01-specification")

    def test_body_protected(self):
        self.assertFalse(self.r.can_write(".claude/skills/sdlc-security/SKILL.md")[0])
        self.assertFalse(self.r.can_write(".claude/agents/security-reviewer.md")[0])


class TestRecords(Base):
    def test_decisions_and_questions(self):
        self.assertEqual(records.add_decision(self.r.dir, "言語は Python", "01-specification", "人の指示"), "DEC-001")
        q1 = records.add_question(self.r.dir, "同時貸出の上限は？", "高", "01-specification")
        records.answer_question(self.r.dir, q1, "3 冊")
        self.assertEqual(records.questions(self.r.dir, only_open=True), [])


class TestIntake(Base):
    def test_md_and_refuse_after_gate(self):
        from intake import convert
        with contextlib.redirect_stdout(io.StringIO()):
            convert.main(["input/req.md"])
        out = (self.r.dir / "docs/01-specification/intake/req.md").read_text(encoding="utf-8")
        self.assertIn("### SRC-001", out)
        import re
        ids = re.findall(r"^### (SRC-\d+)", out, re.MULTILINE)
        self.r.write("docs/01-specification/spec.md", SPEC_DOC.replace("- 追跡元: SRC-001", "- 追跡元: " + ", ".join(ids)))
        self.r.approve("01")
        with self.assertRaises(SystemExit) as cm:
            with contextlib.redirect_stdout(io.StringIO()):
                convert.main(["input/req.md"])
        self.assertIn("承認済み", str(cm.exception.code))


def home_env(home) -> dict:
    return {"HOME": str(home), "USERPROFILE": str(home)}


def install(*args, home=None):
    env = dict(os.environ, **home_env(home or TEST_HOME))
    return subprocess.run([sys.executable, str(INSTALLER), *args], capture_output=True, text=True, encoding="utf-8",
                          errors="replace", env=env)


def run_hook(hook_path, project, tool_input, env=None):
    data = {"tool_input": tool_input, "cwd": str(project)}
    e = dict(os.environ, CLAUDE_PROJECT_DIR=str(project))
    e.pop("SDLC_MAINTENANCE", None)
    e.update(env or {})
    return subprocess.run([sys.executable, str(hook_path)], input=json.dumps(data), capture_output=True, text=True,
                          encoding="utf-8", errors="replace", cwd=project, env=e)


class TestPlacement(unittest.TestCase):
    """~/.claude への配置・プロジェクトの .claude による上書き・プロジェクト単体での配置。"""

    def setUp(self):
        self.home = Path(tempfile.mkdtemp(prefix="sdlc-home-"))
        self.user = self.home / ".claude"
        self.proj = Path(tempfile.mkdtemp(prefix="sdlc-proj-"))
        subprocess.run(["git", "init", "-q"], cwd=self.proj)
        self.env = home_env(self.home)

    def tearDown(self):
        shutil.rmtree(self.home, ignore_errors=True)
        shutil.rmtree(self.proj, ignore_errors=True)

    def sdlc(self, *args):
        return subprocess.run([BASH, str(self.proj / "scripts" / "sdlc"), *args], cwd=self.proj,
                              capture_output=True, text=True, encoding="utf-8", errors="replace",
                              env=dict(os.environ, **self.env))

    def user_init(self):
        r = subprocess.run([sys.executable, str(self.user / "sdlc/bin/sdlc"), "init"], cwd=self.proj,
                           capture_output=True, text=True, encoding="utf-8", errors="replace",
                           env=dict(os.environ, **self.env))
        self.assertEqual(r.returncode, 0, r.stderr)

    def test_user_install_idempotent_and_uninstall(self):
        self.user.mkdir(parents=True)
        (self.user / "settings.json").write_text(json.dumps({"model": "x"}), encoding="utf-8")
        for _ in range(2):
            r = install("--user", home=self.home)
            self.assertEqual(r.returncode, 0, r.stderr)
        for rel in ["sdlc/bin/sdlc", "sdlc/lib/sdlc_lib/testitems.py", "sdlc/RULES.md", "sdlc/phase-flow.md",
                    "hooks/sdlc/guard_write.py", "skills/sdlc/SKILL.md", "skills/phase-acceptance/SKILL.md",
                    "agents/tester.md", "sdlc/templates/test-items.md", "sdlc/checklists/02-design.md",
                    "skills/sdlc-security/SKILL.md", "agents/security-reviewer.md", "sdlc/checklists/security.md",
                    "sdlc/templates/security-review.md"]:
            self.assertTrue((self.user / rel).exists(), rel)
        s = json.loads((self.user / "settings.json").read_text(encoding="utf-8"))
        self.assertEqual(s["model"], "x")
        self.assertEqual(len([h for e in s["hooks"]["PreToolUse"] for h in e["hooks"]]), 3)
        r = install("--user", "--uninstall", home=self.home)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertFalse((self.user / "sdlc").exists())
        self.assertEqual(json.loads((self.user / "settings.json").read_text(encoding="utf-8")), {"model": "x"})

    def test_user_install_project_init_and_gating(self):
        install("--user", home=self.home)
        hook = self.user / "hooks/sdlc/guard_write.py"
        target = {"file_path": str(self.proj / "docs/02-design/a.md")}
        self.assertEqual(run_hook(hook, self.proj, target, self.env).returncode, 0)   # SDLC プロジェクトでなければ何もしない
        self.user_init()
        for rel in ["scripts/sdlc", "scripts/approve", "state/phase.json", "docs/04-acceptance/.gitkeep"]:
            self.assertTrue((self.proj / rel).exists(), rel)
        self.assertIn("@~/.claude/sdlc/RULES.md", (self.proj / "CLAUDE.md").read_text(encoding="utf-8"))
        self.assertIn("workflow: sdlc", (self.proj / "project.yaml").read_text(encoding="utf-8"))
        self.assertIn("配置: user", self.sdlc("where").stdout)
        self.assertEqual(run_hook(hook, self.proj, target, self.env).returncode, 2)

    def test_project_overrides_user(self):
        install("--user", home=self.home)
        r = install("--project", str(self.proj), home=self.home)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.user_init()
        self.assertIn("@.claude/sdlc/RULES.md", (self.proj / "CLAUDE.md").read_text(encoding="utf-8"))
        target = {"file_path": str(self.proj / "docs/02-design/a.md")}
        self.assertEqual(run_hook(self.user / "hooks/sdlc/guard_write.py", self.proj, target, self.env).returncode, 0)
        self.assertEqual(run_hook(self.proj / ".claude/hooks/sdlc/guard_write.py", self.proj, target,
                                  self.env).returncode, 2)
        r = self.sdlc("path", "templates/review.md")
        self.assertEqual(Path(r.stdout.strip()).resolve(), (self.proj / ".claude/sdlc/templates/review.md").resolve())

    def test_source_settings_match_installer(self):
        sys.path.insert(0, str(REPO))
        import importlib
        inst = importlib.import_module("install")
        expected = inst.merged_settings({}, "project")
        actual = json.loads((REPO / ".claude/settings.json").read_text(encoding="utf-8"))
        self.assertEqual(actual, expected)

    def test_body_files_referenced_exist(self):
        """スキル・エージェント・ルールが参照する <SDLC>/ のファイルがすべて存在する。"""
        import re
        body = REPO / ".claude"
        refs = set()
        for p in list(body.rglob("*.md")):
            refs |= set(re.findall(r"<SDLC>/([\w./-]+\.(?:md|yaml))", p.read_text(encoding="utf-8")))
        missing = sorted(r for r in refs if "<" not in r and not (body / "sdlc" / r).exists())
        self.assertEqual(missing, [])


class TestWindowsSupport(Base):
    def test_crlf_does_not_change_gate_hash(self):
        (self.r.dir / "docs/01-specification").mkdir(parents=True)
        (self.r.dir / "docs/01-specification/a.md").write_bytes(b"line1\nline2\n")
        h1, _ = gates.artifacts(self.r.dir, self.r.proj, "01")
        (self.r.dir / "docs/01-specification/a.md").write_bytes(b"line1\r\nline2\r\n")
        self.assertEqual(h1, gates.artifacts(self.r.dir, self.r.proj, "01")[0])

    def test_approve_refuses_inside_claude_code(self):
        r = subprocess.run([sys.executable, str(BIN / "approve"), "01"], capture_output=True, text=True,
                           encoding="utf-8", errors="replace", cwd=self.r.dir, env=dict(os.environ, CLAUDECODE="1"))
        self.assertEqual(r.returncode, 1)
        self.assertIn("Claude Code", r.stderr)

    def test_hook_entry_run_wrapper(self):
        data = json.dumps({"tool_input": {"file_path": str(self.r.dir / "docs/02-design/a.md")},
                           "cwd": str(self.r.dir)})
        env = dict(os.environ, CLAUDE_PROJECT_DIR=str(self.r.dir))
        r = subprocess.run([BASH, str(HOOKS / "run"), "guard_write.py"], input=data, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", env=env)
        self.assertEqual(r.returncode, 2, r.stderr)
        r = subprocess.run([BASH, str(HOOKS / "run"), "guard_write.py"], input=data, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", env=dict(env, SDLC_PYTHON="/nonexistent/python"))
        self.assertEqual(r.returncode, 2)
        self.assertIn("Python", r.stderr)


if __name__ == "__main__":
    unittest.main()
