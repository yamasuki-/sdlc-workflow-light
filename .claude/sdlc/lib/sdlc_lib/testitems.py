"""テスト項目（TC）と受入確認の結果シート（docs/04-acceptance/results.md）。

- TC は工程 01 で仕様（SPEC / NFR）から作る「何を見れば OK か」のリスト（docs/ 内の `### TC-xxx`）。
- 結果シートは Markdown の表。人が「結果」列に OK / NG / 保留 / 対象外 を記入する（Claude は記入しない）。
  シートの作成・更新（sheet）、再確認のための差し戻し（reset）、Excel との出力・取り込み（export / import）は CLI で行う。
- 対象は種別「手動」「両方」の TC。種別「自動」の TC はテストコード（`trace: TC-xxx`）と commands.test で確認する。
"""
from __future__ import annotations

import re
from pathlib import Path

from . import trace
from .config import RESULTS_FILE
from .state import now

RESULT_VALUES = ("OK", "NG", "保留", "対象外", "未実施")
PASS_VALUES = ("OK", "対象外")
COLUMNS = ["TC", "タイトル", "種別", "結果", "確認者", "日付", "メモ"]
FIELD_RE = re.compile(r"^\s*[-*]\s*([^:：\s][^:：]*?)\s*[:：]\s*(.*)$")
ROW_RE = re.compile(r"^\|\s*(TC-\d{3,})\s*\|")
# Excel の列（出力・取り込み）
XLSX_COLUMNS = ["TC", "タイトル", "追跡元", "種別", "前提", "手順", "確認ポイント", "期待結果",
                "結果", "確認者", "日付", "メモ"]
XLSX_INPUT = ("結果", "確認者", "日付", "メモ")

SHEET_HEAD = """# 受入確認 結果シート

<!--
人が「結果」列に OK / NG / 保留 / 対象外 を記入する（Claude は記入しない）。
- NG: メモに現象を書く（Claude が不具合として起票し /defect で対応する）
- 対象外: メモに理由を書く
- 確認者・日付も記入する
TC の追加・削除は `scripts/sdlc testitems sheet` で反映する（記入済みの結果は残る）。
Excel で記入する場合は `scripts/sdlc testitems export` → 記入 → `scripts/sdlc testitems import <ファイル>`。
-->

"""


class ItemsError(Exception):
    pass


def results_path(root) -> Path:
    return Path(root) / RESULTS_FILE


def _clean(s: str) -> str:
    return str(s or "").replace("|", "／").replace("\r", "").replace("\n", "<br>").strip()


def details(root, g: trace.Graph) -> dict[str, dict]:
    """TC ごとの記載内容（タイトル・追跡元・種別・前提・手順・確認ポイント・期待結果）。"""
    out = {}
    cache: dict[str, list[str]] = {}
    for tc, d in g.defs.items():
        if trace.kind_of(tc) != "TC":
            continue
        lines = cache.setdefault(d["file"], (Path(root) / d["file"]).read_text(encoding="utf-8").splitlines())
        title = re.sub(r"^" + re.escape(tc) + r"\s*[:：]?\s*", "", d["title"]).strip()
        info = {"TC": tc, "タイトル": title, "種別": d.get("kind") or "", "file": d["file"]}
        key = None
        for line in lines[d["line"]:]:
            if line.startswith("#"):
                break
            m = FIELD_RE.match(line)
            if m and not line.startswith((" ", "\t")):
                key, val = m.group(1), m.group(2).strip()
                info[key] = val
            elif key and line.strip():
                info[key] = (info.get(key, "") + "\n" + line.strip()).strip()
        out[tc] = info
    return out


def manual_tcs(g: trace.Graph) -> list[str]:
    def num(t):
        return int(t.split("-")[1])
    return sorted((t for t, d in g.defs.items() if trace.kind_of(t) == "TC" and d.get("kind") in trace.MANUAL_KINDS
                   and not d.get("exempt")), key=num)


# ---------------------------------------------------------------- 結果シート

def read_results(root) -> dict[str, dict]:
    p = results_path(root)
    out = {}
    if not p.exists():
        return out
    for line in p.read_text(encoding="utf-8").splitlines():
        if not ROW_RE.match(line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        cells += [""] * (len(COLUMNS) - len(cells))
        row = dict(zip(COLUMNS, cells))
        out[row["TC"]] = {"title": row["タイトル"], "kind": row["種別"], "result": row["結果"] or "未実施",
                          "by": row["確認者"], "date": row["日付"], "memo": row["メモ"]}
    return out


def render_results(rows: list[dict]) -> str:
    lines = [SHEET_HEAD.rstrip("\n"), "", "| " + " | ".join(COLUMNS) + " |", "|" + "---|" * len(COLUMNS)]
    for r in rows:
        vals = [r["tc"], r["title"], r["kind"], r["result"], r["by"], r["date"], r["memo"]]
        lines.append("| " + " | ".join(_clean(v) for v in vals) + " |")
    return "\n".join(lines) + "\n"


def write_results(root, rows: list[dict]) -> None:
    p = results_path(root)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(render_results(rows), encoding="utf-8")


def synced_rows(root, g: trace.Graph) -> tuple[list[dict], list[str], list[str]]:
    """TC 一覧に合わせた結果シートの行（記入済みの結果は残し、追加は未実施、削除された TC は外す）。書き込まない。"""
    cur = read_results(root)
    det = details(root, g)
    rows, added = [], []
    for tc in manual_tcs(g):
        old = cur.get(tc)
        if old is None:
            added.append(tc)
            old = {"result": "未実施", "by": "", "date": "", "memo": ""}
        rows.append({"tc": tc, "title": det[tc]["タイトル"], "kind": det[tc]["種別"], **{k: old[k] for k in
                     ("result", "by", "date", "memo")}})
    removed = sorted(set(cur) - {r["tc"] for r in rows})
    return rows, added, removed


def sync_sheet(root, g: trace.Graph) -> dict:
    rows, added, removed = synced_rows(root, g)
    write_results(root, rows)
    return {"rows": len(rows), "added": added, "removed": removed}


def reset(root, g: trace.Graph, tcs: list[str], reason: str) -> list[str]:
    """再確認が必要な TC の結果を未実施に戻す（不具合修正・仕様変更の後）。"""
    unknown = [t for t in tcs if t not in g.defs]
    if unknown:
        raise ItemsError("未定義の TC: " + ", ".join(unknown))
    done = []
    rows = []
    for r in synced_rows(root, g)[0]:
        tc = r.pop("tc")
        if tc in tcs:
            memo = f"再確認（{reason}）" + (f" 前回: {r['result']} {r['memo']}".rstrip() if r["result"] != "未実施" else "")
            r = {**r, "result": "未実施", "by": "", "date": "", "memo": memo}
            done.append(tc)
        rows.append({"tc": tc, **r})
    write_results(root, rows)
    return done


def problems(root, g: trace.Graph) -> list[str]:
    """受入確認の完了条件: 手動確認の TC がすべて OK（または理由付きの対象外）。"""
    cur = read_results(root)
    if not results_path(root).exists():
        return [f"{RESULTS_FILE} がありません（`scripts/sdlc testitems sheet` で作成し、人が記入する）"]
    out = []
    for tc in manual_tcs(g):
        r = cur.get(tc)
        if r is None:
            out.append(f"{tc}: 結果シートにありません（`scripts/sdlc testitems sheet` で反映する）")
        elif r["result"] not in RESULT_VALUES:
            out.append(f"{tc}: 結果 {r['result']!r} は {'/'.join(RESULT_VALUES[:4])} のいずれかで記入する")
        elif r["result"] == "対象外" and not r["memo"]:
            out.append(f"{tc}: 対象外の理由がメモにありません")
        elif r["result"] not in PASS_VALUES:
            out.append(f"{tc}: {r['result']}" + (f"（{r['memo']}）" if r["memo"] else ""))
    return out


def summary(root, g: trace.Graph) -> dict:
    cur = read_results(root)
    c = {v: 0 for v in RESULT_VALUES}
    for tc in manual_tcs(g):
        res = (cur.get(tc) or {}).get("result", "未実施")
        c[res if res in c else "未実施"] += 1
    autos = [t for t, d in g.defs.items() if trace.kind_of(t) == "TC" and d.get("kind") in trace.AUTO_KINDS]
    c["自動（テストコードあり）"] = sum(1 for t in autos if trace.automated(g, t))
    c["自動（テストコードなし）"] = sum(1 for t in autos if not trace.automated(g, t))
    return c


def report(root, g: trace.Graph) -> str:
    """仕様ごとのテスト項目と確認状況。"""
    cur = read_results(root)
    det = details(root, g)
    lines = [f"# 受入確認レポート", "", f"- 作成日時: {now()}", ""]
    s = summary(root, g)
    lines += ["## 集計", "", "| 区分 | 件数 |", "|---|---|"] + [f"| {k} | {v} |" for k, v in s.items()]
    lines += ["", "## 仕様ごとの確認状況", "", "| 仕様 | TC | 種別 | 手動確認 | 自動テスト |", "|---|---|---|---|---|"]
    specs = sorted(i for i in g.defs if trace.kind_of(i) in ("SPEC", "NFR"))
    for sp in specs:
        tcs = sorted(c for c in g.children.get(sp, ()) if trace.kind_of(c) == "TC")
        if not tcs:
            lines.append(f"| {sp} | — | | | |")
        for tc in tcs:
            kind = det.get(tc, {}).get("種別", "")
            man = (cur.get(tc) or {}).get("result", "未実施") if kind in trace.MANUAL_KINDS else "—"
            auto = "<br>".join(trace.automated(g, tc)) or ("なし" if kind in trace.AUTO_KINDS else "—")
            lines.append(f"| {sp} | {tc} | {kind} | {man} | {auto} |")
    ng = [(t, r) for t, r in cur.items() if r["result"] in ("NG", "保留")]
    if ng:
        lines += ["", "## NG・保留", ""] + [f"- {t} [{r['result']}] {r['memo']}" for t, r in ng]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- Excel（任意）

def _openpyxl():
    try:
        import openpyxl
    except ImportError as e:
        raise ItemsError("Excel の出力・取り込みには openpyxl が必要です（pip install -r requirements.txt）。") from e
    return openpyxl


def export_xlsx(root, g: trace.Graph, out: Path) -> int:
    """手動確認の TC と現在の結果を Excel に出力する。人は「結果」〜「メモ」列だけを記入する。"""
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    det = details(root, g)
    cur = read_results(root)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "受入確認"
    ws.append(XLSX_COLUMNS)
    head = PatternFill("solid", fgColor="DDDDDD")
    inp = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head
    tcs = manual_tcs(g)
    for tc in tcs:
        d, r = det[tc], cur.get(tc) or {}
        res = r.get("result", "未実施")
        ws.append([tc, d["タイトル"], d.get("追跡元", ""), d["種別"], d.get("前提", ""), d.get("手順", ""),
                   d.get("確認ポイント", ""), d.get("期待結果", ""), "" if res == "未実施" else res,
                   r.get("by", ""), r.get("date", ""), str(r.get("memo", "")).replace("<br>", "\n")])
    widths = [10, 30, 14, 8, 30, 50, 40, 40, 10, 12, 12, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if XLSX_COLUMNS[c.column - 1] in XLSX_INPUT:
                c.fill = inp
    if tcs:
        col = openpyxl.utils.get_column_letter(XLSX_COLUMNS.index("結果") + 1)
        dv = DataValidation(type="list", formula1='"OK,NG,保留,対象外"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(tcs) + 1}")
    ws.freeze_panes = "B2"
    note = wb.create_sheet("記入方法")
    for line in ("黄色の列（結果・確認者・日付・メモ）だけを記入してください。",
                 "結果: OK / NG / 保留 / 対象外（空欄は未実施）",
                 "NG はメモに現象を、対象外はメモに理由を書いてください。",
                 "記入後: scripts/sdlc testitems import <このファイル> で結果シートに取り込みます。"):
        note.append([line])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return len(tcs)


def import_xlsx(root, g: trace.Graph, src: Path) -> dict:
    """Excel に記入された結果を結果シートに取り込む（結果・確認者・日付・メモだけを読む）。"""
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["受入確認"] if "受入確認" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ItemsError("シートが空です。")
    head = [str(c or "").strip() for c in rows[0]]
    missing = [c for c in ("TC",) + XLSX_INPUT if c not in head]
    if missing:
        raise ItemsError("列が見つかりません: " + ", ".join(missing))
    idx = {c: head.index(c) for c in ("TC",) + XLSX_INPUT}
    cur = {r.pop("tc"): r for r in synced_rows(root, g)[0]}
    errors, updated = [], []
    for n, row in enumerate(rows[1:], 2):
        tc = str(row[idx["TC"]] or "").strip()
        if not tc:
            continue
        if tc not in cur:
            errors.append(f"{n} 行目: {tc} は手動確認の TC ではありません")
            continue
        res = str(row[idx["結果"]] or "").strip() or "未実施"
        if res not in RESULT_VALUES:
            errors.append(f"{n} 行目: {tc} の結果 {res!r} は {'/'.join(RESULT_VALUES[:4])} のいずれか")
            continue
        date = row[idx["日付"]]
        date = date.strftime("%Y-%m-%d") if hasattr(date, "strftime") else str(date or "").strip()
        new = {"result": res, "by": str(row[idx["確認者"]] or "").strip(), "date": date,
               "memo": str(row[idx["メモ"]] or "").strip()}
        old = cur[tc]
        if any(str(old[k]) != _clean(new[k]) for k in new):
            updated.append(tc)
        cur[tc] = {**old, **new}
    if errors:
        raise ItemsError("取り込めません（結果シートは変更していません）:\n  - " + "\n  - ".join(errors))
    write_results(root, [{"tc": tc, **r} for tc, r in cur.items()])
    return {"updated": updated, "rows": len(cur)}
