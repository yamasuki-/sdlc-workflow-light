"""Excel: 出典 = シート名＋セル番地。結合セルは値を展開し、表形式の要求一覧を行単位の断片にする。"""
from __future__ import annotations

from pathlib import Path

from .base import HIGH, LOW, MID, Fragment


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def convert(path: Path, assets_dir: Path, rel_assets: str) -> list[Fragment]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path, data_only=True)
    out: list[Fragment] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            out.append(Fragment(f"{ws.title}（非表示シート）", "注記", "非表示シートです。対象に含めるか確認してください。", LOW))
        # 結合セルの展開: 左上の値を範囲全体に複製する
        grid: dict[tuple[int, int], str] = {}
        merged_cells: set[tuple[int, int]] = set()
        for row in ws.iter_rows():
            for c in row:
                v = _s(c.value)
                if v:
                    grid[(c.row, c.column)] = v
        for rng in ws.merged_cells.ranges:
            v = grid.get((rng.min_row, rng.min_col), "")
            for r in range(rng.min_row, rng.max_row + 1):
                for col in range(rng.min_col, rng.max_col + 1):
                    merged_cells.add((r, col))
                    if v:
                        grid[(r, col)] = v
        if not grid:
            continue
        rows = sorted({r for r, _ in grid})
        cols = sorted({c for _, c in grid})
        header_row = next((r for r in rows if sum(1 for c in cols if (r, c) in grid) >= 2), None)
        header = {c: grid.get((header_row, c), get_column_letter(c)) for c in cols} if header_row else {}
        for r in rows:
            if r == header_row:
                continue
            cells = [(c, grid[(r, c)]) for c in cols if (r, c) in grid]
            if not cells:
                continue
            a, b = cells[0][0], cells[-1][0]
            src = f"{ws.title}!{get_column_letter(a)}{r}" + (f":{get_column_letter(b)}{r}" if b != a else "")
            if header_row and r > header_row:
                text = "\n".join(f"{header.get(c, get_column_letter(c))}: {v}" for c, v in cells)
                kind = "表行"
            else:
                text = "\n".join(v for _, v in cells)
                kind = "セル"
            conf, notes = HIGH, []
            if any((r, c) in merged_cells for c, _ in cells):
                conf, notes = MID, ["結合セルを展開しています（値の適用範囲を確認）"]
            out.append(Fragment(src, kind, text, conf, notes))
        # 埋め込み画像
        for i, img in enumerate(getattr(ws, "_images", []), 1):
            try:
                data = img._data()
                fmt = (getattr(img, "format", None) or "png").lower()
                name = f"{path.stem}-{ws.title}-img{i}.{fmt}"
                assets_dir.mkdir(parents=True, exist_ok=True)
                (assets_dir / name).write_bytes(data)
                anchor = img.anchor._from
                cell = f"{get_column_letter(anchor.col + 1)}{anchor.row + 1}"
                out.append(Fragment(f"{ws.title}!{cell}（画像）", "図", "要画像読取", LOW,
                                    ["画像の内容を読み取り、本文と確信度を記入すること"], [f"{rel_assets}/{name}"]))
            except Exception as e:  # noqa: BLE001
                out.append(Fragment(f"{ws.title}（画像{i}）", "図", f"画像を抽出できませんでした: {e}", LOW))
    return out
