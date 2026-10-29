"""PowerPoint: 出典 = スライド番号＋図形。図・構成図・権限表は画像化して読み取り対象にし、確信度を付与する。

LibreOffice（soffice）があればスライド全体を PNG 化し、構成図らしきスライドの読み取りに使う。
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from .base import HIGH, LOW, MID, Fragment


def _render_slides(path: Path, assets_dir: Path) -> dict[int, str]:
    """soffice で PDF 化 → pdftoppm で PNG 化。どちらかが無ければ空を返す。"""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        return {}
    try:
        with tempfile.TemporaryDirectory() as td:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, str(path)],
                           capture_output=True, timeout=180, check=True)
            pdf = Path(td) / (path.stem + ".pdf")
            prefix = Path(td) / "s"
            subprocess.run([pdftoppm, "-png", "-r", "80", str(pdf), str(prefix)], capture_output=True, timeout=180, check=True)
            out = {}
            assets_dir.mkdir(parents=True, exist_ok=True)
            for png in sorted(Path(td).glob("s-*.png")):
                n = int(png.stem.split("-")[-1])
                dst = assets_dir / f"{path.stem}-slide{n}.png"
                shutil.copy(png, dst)
                out[n] = dst.name
            return out
    except (OSError, subprocess.SubprocessError, ValueError):
        return {}


def convert(path: Path, assets_dir: Path, rel_assets: str) -> list[Fragment]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    out: list[Fragment] = []
    diagram_slides: list[int] = []

    for sn, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        base = f"スライド{sn}" + (f"「{title}」" if title else "")
        visual = {"pic": 0, "connector": 0, "auto": 0}

        def walk(shapes, group=""):
            for sh in shapes:
                name = f"{group}{sh.name}"
                src = f"{base} / 図形「{name}」"
                st = sh.shape_type
                if st == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes, group=f"{name} > ")
                    continue
                if st == MSO_SHAPE_TYPE.PICTURE:
                    visual["pic"] += 1
                    img = sh.image
                    fname = f"{path.stem}-slide{sn}-{sh.shape_id}.{img.ext}"
                    assets_dir.mkdir(parents=True, exist_ok=True)
                    (assets_dir / fname).write_bytes(img.blob)
                    out.append(Fragment(src, "図", "要画像読取", LOW,
                                        ["画像の内容を読み取り、本文と確信度を記入すること"], [f"{rel_assets}/{fname}"]))
                    continue
                if getattr(sh, "is_connector", False) or st == MSO_SHAPE_TYPE.LINE:
                    visual["connector"] += 1
                    continue
                if getattr(sh, "has_table", False) and sh.has_table:
                    rows = list(sh.table.rows)
                    header = [c.text.strip() for c in rows[0].cells] if rows else []
                    for rn, row in enumerate(rows[1:], 2):
                        cells = [c.text.strip() for c in row.cells]
                        if not any(cells):
                            continue
                        text = "\n".join(f"{h or f'列{i + 1}'}: {v}" for i, (h, v) in enumerate(zip(header, cells)) if v)
                        marks = {v for v in cells if v in ("○", "◯", "×", "✓", "-", "△")}
                        notes = ["記号（○×等）の意味を凡例で確認"] if marks else []
                        out.append(Fragment(f"{src} 行{rn}", "表行", text, MID if marks else HIGH, notes))
                    continue
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    ch = sh.chart
                    try:
                        cats = list(ch.plots[0].categories)
                        lines = [f"カテゴリ: {', '.join(map(str, cats))}"]
                        for s in ch.series:
                            lines.append(f"{s.name}: {', '.join(map(str, s.values))}")
                        out.append(Fragment(src, "グラフ", "\n".join(lines), MID, ["グラフの意図（要求か参考値か）を確認"]))
                    except Exception:  # noqa: BLE001
                        out.append(Fragment(src, "グラフ", "グラフのデータを抽出できませんでした", LOW))
                    continue
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    text = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip()).strip()
                    if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                        visual["auto"] += 1
                    if not text or sh == slide.shapes.title:
                        continue
                    out.append(Fragment(src, "テキスト", text, HIGH))

        walk(slide.shapes)
        # 図形とコネクタで描かれた構成図: テキストを拾っても関係（矢印）が失われるため画像で読む
        if visual["connector"] >= 1 or visual["auto"] >= 4:
            diagram_slides.append(sn)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                out.append(Fragment(f"{base} / ノート", "ノート", notes, MID, ["発表者ノート。要求か補足かを確認"]))

    if diagram_slides:
        rendered = _render_slides(path, assets_dir)
        for sn in diagram_slides:
            name = rendered.get(sn)
            note = "図形・矢印で構成された図（構成図・フロー等）。スライド画像から関係を読み取り、本文と確信度を記入すること"
            if name:
                out.append(Fragment(f"スライド{sn} / スライド全体（画像化）", "図", "要画像読取", LOW, [note],
                                    [f"{rel_assets}/{name}"]))
            else:
                out.append(Fragment(f"スライド{sn} / スライド全体", "図", "要画像読取（画像化できず: soffice / pdftoppm が無い）",
                                    LOW, [note + "。人に画像を提供してもらうか質問リストに回す"]))
    return out
